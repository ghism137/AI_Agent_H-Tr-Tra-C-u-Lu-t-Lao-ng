from io import BytesIO
from pathlib import Path

from docx import Document

from backend.ingestion.docx_extractor import extract_docx
from backend.ingestion.parser import parse_legal_document


def test_merged_table_keeps_cell_and_row_locators():
    document = Document()
    document.add_paragraph("Điều 1. Mức áp dụng")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Vùng lương"
    table.cell(1, 0).text = "I"
    table.cell(1, 1).text = "5.000.000 đồng"
    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)

    blocks = extract_docx(buffer)
    source_table = next(block for block in blocks if block["kind"] == "table")
    assert source_table["cells"][0]["column_span"] == 2
    assert source_table["cells"][0]["locator"].endswith("/row:1/col:1")

    parsed = parse_legal_document("", {"doc_number": "example"}, blocks)
    article = parsed["articles"][0]
    assert "Vùng lương" in article["content"]
    assert "5.000.000 đồng" in article["content"]
    assert any(seg["locator"].endswith("/row:2") for seg in article["content_segments"])
    assert parsed["source_tables"][0]["cells"][0]["column_span"] == 2


def test_form_12_vertical_merge_does_not_repeat_header_text():
    source = Path(__file__).resolve().parents[1] / "data/raw/word/2025_979 + 980_188-2025-NĐ-CP.docx"
    table = next(block for block in extract_docx(source) if block["locator"] == "body/table:374")
    assert table["rows"][0][0] == "STT"
    assert table["rows"][1] == ["", "", "", ""]
    assert all(cell["vertical_merge"] == "continue" for cell in table["cells"] if cell["row"] == 2)
